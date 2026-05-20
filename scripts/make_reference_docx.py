# man_hours: 4.0
"""Build the pandoc reference.docx template from report_format.json.

Do not edit ``reference.docx`` by hand. It is a generated Pandoc style cache;
``report_format.json`` is the sole layout source of truth. Callers should use
``ensure_reference_docx()`` so the template is rebuilt when the JSON changes.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from report_format_config import ReportFormatConfig


def _apply_font(style, *, name: str, size_pt: float, bold: bool = False,
                italic: bool = False) -> None:
    font = style.font
    font.name = name
    font.size = Pt(size_pt)
    font.bold = bold
    font.italic = italic
    rpr = style.element.get_or_add_rPr()
    for tag in ("rFonts",):
        existing = rpr.find(qn(f"w:{tag}"))
        if existing is not None:
            rpr.remove(existing)
    rfonts = OxmlElement("w:rFonts")
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:cs"), name)
    rfonts.set(qn("w:eastAsia"), name)
    rpr.append(rfonts)


def _apply_paragraph_format(style, *, alignment, space_before_pt=0,
                             space_after_pt=0, line_spacing=1.0,
                             page_break_before: bool = False,
                             keep_with_next: bool = False) -> None:
    pf = style.paragraph_format
    pf.alignment = alignment
    pf.space_before = Pt(space_before_pt)
    pf.space_after = Pt(space_after_pt)
    pf.line_spacing = line_spacing
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.page_break_before = page_break_before
    pf.keep_with_next = keep_with_next


def _set_a4_section(section, config: ReportFormatConfig) -> None:
    margins = config.margins_mm()
    section.page_height = Cm(29.7)
    section.page_width = Cm(21.0)
    section.top_margin = Cm(margins["top"] / 10)
    section.bottom_margin = Cm(margins["bottom"] / 10)
    section.left_margin = Cm(margins["inside"] / 10)
    section.right_margin = Cm(margins["outside"] / 10)
    page = config.data["page"]
    section.header_distance = Cm(page["running_header_height_mm"] / 10)
    section.footer_distance = Cm(page["footer_height_mm"] / 10)


def _set_footer_page_number(section, config: ReportFormatConfig) -> None:
    footer = config.data["typography"]["footer_page_number"]
    footer_paragraph = section.footer.paragraphs[0]
    footer_paragraph.alignment = config.alignment(str(footer["alignment"]))
    footer_paragraph.clear()
    run = footer_paragraph.add_run()
    run.font.name = str(footer["font_family"])
    run.font.size = Pt(float(footer["size_pt"]))

    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    run._r.append(begin)
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    run._r.append(instr_text)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(end)


def _configure_heading(styles, config: ReportFormatConfig, level: int) -> None:
    spec = config.data["typography"]["headings"][f"h{level}"]
    style = styles[f"Heading {level}"]
    weight = str(spec.get("weight", "regular"))
    _apply_font(
        style,
        name=str(spec["font_family"]),
        size_pt=float(spec["size_pt"]),
        bold=weight in ("bold", "semibold"),
    )
    _apply_paragraph_format(
        style,
        alignment=config.alignment(str(spec["alignment"])),
        space_before_pt=18 if level == 1 else 12,
        space_after_pt=6,
        line_spacing=config.line_spacing("body"),
        page_break_before=level == 1,
        keep_with_next=True,
    )


def build_reference_docx(
    path: Path | None = None,
    config: ReportFormatConfig | None = None,
) -> Path:
    fmt = config or ReportFormatConfig.load()
    out_path = path or fmt.build_path("reference_docx")
    out_path.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()
    _set_a4_section(doc.sections[0], fmt)
    _set_footer_page_number(doc.sections[0], fmt)

    styles = doc.styles
    body = fmt.data["typography"]["body"]
    title = fmt.data["typography"]["title"]
    subtitle = fmt.data["typography"]["subtitle"]

    normal = styles["Normal"]
    _apply_font(normal, name=str(body["font_family"]), size_pt=float(body["size_pt"]))
    _apply_paragraph_format(
        normal,
        alignment=fmt.alignment(str(body["alignment"])),
        space_after_pt=6,
        line_spacing=float(body["line_spacing"]),
    )

    title_style = styles["Title"]
    _apply_font(
        title_style,
        name=str(title["font_family"]),
        size_pt=float(title["size_pt"]),
        bold=str(title.get("weight")) == "bold",
    )
    _apply_paragraph_format(
        title_style,
        alignment=fmt.alignment(str(title["alignment"])),
        space_before_pt=24,
        space_after_pt=12,
        line_spacing=float(body["line_spacing"]),
        keep_with_next=True,
    )

    if "Subtitle" in [s.name for s in styles]:
        subtitle_style = styles["Subtitle"]
        _apply_font(
            subtitle_style,
            name=str(subtitle["font_family"]),
            size_pt=float(subtitle["size_pt"]),
        )
        _apply_paragraph_format(
            subtitle_style,
            alignment=fmt.alignment(str(subtitle["alignment"])),
            space_after_pt=12,
            line_spacing=float(body["line_spacing"]),
            keep_with_next=True,
        )

    for level in range(1, 5):
        _configure_heading(styles, fmt, level)

    if "Source Code" in [s.name for s in styles]:
        source_code = styles["Source Code"]
        _apply_font(source_code, name="Consolas", size_pt=10)
        _apply_paragraph_format(
            source_code,
            alignment=WD_ALIGN_PARAGRAPH.LEFT,
            space_after_pt=6,
            line_spacing=float(body["line_spacing"]),
        )

    doc.save(out_path)
    return out_path


if __name__ == "__main__":
    written = build_reference_docx()
    print(f"Wrote reference template: {written}")
