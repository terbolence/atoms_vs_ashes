# man_hours: 1.2
"""Build the v1.2 landscape results-table deliverable.

This module is the only supported build path for ``atoms_vs_ashes_results_table``
(``.md``, ``.csv``, and landscape ``.docx``). It always regenerates Markdown from
country ledgers before exporting Word.

Run directly::

    python scripts/build_results_table_deliverable.py

Or through the report build::

    python scripts/build_report.py --side-deliverables-only
    python scripts/build_report.py   # unless --skip-side-deliverables

Do not use ``scripts/export_markdown_docx.py`` on the results-table Markdown.
"""

from __future__ import annotations

import argparse
import shutil
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm

from report_docx_postprocess import postprocess_docx
from report_format_config import (
    DEFAULT_FORMAT_PATH,
    ReportFormatConfig,
    ensure_reference_docx,
)
from results_table_data import build_results_markdown
from results_table_model import RESULTS_TABLE_OUTPUT_STEM, configure_from_format

OUTPUT_STEM = RESULTS_TABLE_OUTPUT_STEM

_MANUAL_EXPORT_ERROR = (
    "The results table must be built with scripts/build_results_table_deliverable.py "
    "(or python scripts/build_report.py --side-deliverables-only), which regenerates "
    "Markdown and CSV from country ledgers before exporting DOCX. "
    "export_markdown_docx.py is not supported for this deliverable."
)


def is_results_table_markdown(path: Path) -> bool:
    return path.resolve().stem == OUTPUT_STEM


def reject_manual_results_table_export(path: Path) -> None:
    if is_results_table_markdown(path):
        raise ValueError(_MANUAL_EXPORT_ERROR)


def _run_pandoc(markdown_path: Path, output_docx: Path, reference_docx: Path) -> None:
    if shutil.which("pandoc") is None:
        raise RuntimeError("pandoc not found on PATH; install pandoc first.")
    markdown_path = markdown_path.resolve()
    output_docx = output_docx.resolve()
    reference_docx = reference_docx.resolve()
    command = [
        "pandoc",
        markdown_path.name,
        "-o",
        str(output_docx),
        "--from=markdown+pipe_tables+header_attributes+fenced_code_blocks",
        "--to=docx",
        "--standalone",
        f"--reference-doc={reference_docx}",
        "--metadata",
        "title=Atoms vs Ashes Results Table",
    ]
    subprocess.run(command, check=True, cwd=markdown_path.parent)


def _set_landscape_a3(docx_path: Path, config: ReportFormatConfig) -> None:
    doc = Document(docx_path)
    margins = config.margins_mm()
    for section in doc.sections:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width = Mm(420)
        section.page_height = Mm(297)
        section.top_margin = Mm(margins["top"])
        section.bottom_margin = Mm(margins["bottom"])
        section.left_margin = Mm(margins["inside"])
        section.right_margin = Mm(margins["outside"])
    doc.save(docx_path)


def _a3_landscape_content_emu(config: ReportFormatConfig) -> tuple[int, int]:
    """Printable width/height in EMU for a map-only A3 landscape page.

    Reserves ~14 mm below the top margin for the country H2 line so the
    PNG can use the remaining sheet without clipping.
    """
    margins = config.margins_mm()
    width_mm = 420.0 - margins["inside"] - margins["outside"]
    height_mm = 297.0 - margins["top"] - margins["bottom"] - 14.0
    return int(Mm(width_mm).emu), int(Mm(height_mm).emu)


def _paragraph_is_site_status_map(paragraph) -> bool:
    """True when the paragraph embeds a country site-status map figure."""
    for run in paragraph.runs:
        for drawing in run._element.findall(".//" + qn("w:drawing")):
            for descr in drawing.findall(".//" + qn("wp:docPr")):
                text = (descr.get("descr") or descr.get("title") or "").lower()
                if "site status map" in text:
                    return True
    return False


def _resize_inline_images(paragraph, *, max_cx: int, max_cy: int) -> None:
    """Fit inline pictures to the map page content box (upscale or
    downscale, preserving aspect ratio).

    The source PNG ships at 4800x3150 px (16x10.5 in at 300 dpi), so
    upscaling by Pandoc's default 6.48 in width to the full A3
    landscape printable area does not introduce pixelation.

    Also scales the corresponding ``a:ext`` inside ``wp:inline``
    (extent of the contained graphicFrame) so Word honors the new
    size.
    """
    for run in paragraph.runs:
        for drawing in run._element.findall(".//" + qn("w:drawing")):
            for inline in drawing.findall(".//" + qn("wp:inline")):
                extent = inline.find(qn("wp:extent"))
                if extent is None:
                    continue
                cx = int(extent.get("cx") or 0)
                cy = int(extent.get("cy") or 0)
                if cx <= 0 or cy <= 0:
                    continue
                scale = min(max_cx / cx, max_cy / cy)
                new_cx = int(cx * scale)
                new_cy = int(cy * scale)
                extent.set("cx", str(new_cx))
                extent.set("cy", str(new_cy))
                a_ns = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
                pic_ns = "{http://schemas.openxmlformats.org/drawingml/2006/picture}"
                for ext in inline.findall(
                    f".//{pic_ns}spPr/{a_ns}xfrm/{a_ns}ext"
                ):
                    ext.set("cx", str(new_cx))
                    ext.set("cy", str(new_cy))


def _is_pagebreak_only_paragraph(p_elem) -> bool:
    """True when ``p_elem`` is a Pandoc-style empty page-break paragraph.

    Matches ``<w:p>`` whose only run contents are ``<w:br w:type="page"/>``
    elements with no ``<w:t>`` text, no ``<w:drawing>`` (image), and no
    ``<w:tab>`` siblings. ``<w:rPr>`` (run properties) and ``<w:pPr>``
    (paragraph properties) are ignored.
    """
    has_break = False
    for child in p_elem:
        tag = child.tag
        if tag == qn("w:pPr"):
            continue
        if tag != qn("w:r"):
            return False
        for run_child in child:
            rtag = run_child.tag
            if rtag == qn("w:rPr"):
                continue
            if rtag == qn("w:br") and run_child.get(qn("w:type")) == "page":
                has_break = True
                continue
            return False
    return has_break


def _set_page_break_before(p_elem) -> None:
    """Idempotently add ``<w:pageBreakBefore/>`` to the paragraph's ``<w:pPr>``."""
    pPr = p_elem.find(qn("w:pPr"))
    if pPr is None:
        pPr = OxmlElement("w:pPr")
        p_elem.insert(0, pPr)
    if pPr.find(qn("w:pageBreakBefore")) is None:
        pPr.append(OxmlElement("w:pageBreakBefore"))


def _strip_pagebreak_only_paragraphs(doc) -> int:
    """Replace every empty ``<w:br type=page>`` paragraph with a
    ``<w:pageBreakBefore/>`` on the next paragraph.

    Pandoc emits each results-table page break as its own empty
    ``<w:p>``. Word renders that empty paragraph as a leading line at
    the top of the new page, which (combined with the 239 mm A3
    landscape map image plus ``space_before/after``) overflows the
    printable area and forces ``keep_together`` to push the entire map
    onto the next page, leaving a blank page behind. Migrating the
    break property onto the following paragraph removes the leading
    line and the orphan-page side effect.
    """
    body = doc.element.body
    nodes = list(body)
    removed = 0
    for index, node in enumerate(nodes):
        if node.tag != qn("w:p"):
            continue
        if not _is_pagebreak_only_paragraph(node):
            continue
        for follower in nodes[index + 1:]:
            if follower.tag == qn("w:p"):
                _set_page_break_before(follower)
                break
        body.remove(node)
        removed += 1
    return removed


def _postprocess_map_pages(
    docx_path: Path, config: ReportFormatConfig,
) -> dict[str, int]:
    """Fit each country status map to the A3 landscape page.

    Round 3 of the layout polish (2026-05-26): also strip every empty
    Pandoc page-break paragraph and migrate the break onto the next
    paragraph (`<w:pageBreakBefore/>`). This eliminates the parasitic
    leading line that was forcing maps onto the next page and leaving
    empty pages behind.
    """
    doc = Document(docx_path)
    max_cx, max_cy = _a3_landscape_content_emu(config)
    stripped = _strip_pagebreak_only_paragraphs(doc)
    map_pages = 0
    for paragraph in doc.paragraphs:
        if not _paragraph_is_site_status_map(paragraph):
            continue
        map_pages += 1
        pf = paragraph.paragraph_format
        pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf.space_before = Mm(2)
        pf.space_after = Mm(2)
        pf.keep_with_next = False
        pf.keep_together = False
        _resize_inline_images(paragraph, max_cx=max_cx, max_cy=max_cy)
    doc.save(docx_path)
    return {"map_pages": map_pages, "pagebreaks_stripped": stripped}


def build_results_table_deliverable(
    *,
    format_path: Path = DEFAULT_FORMAT_PATH,
    output_dir: Path | None = None,
) -> dict[str, int | str]:
    fmt = ReportFormatConfig.load(format_path)
    configure_from_format(fmt)
    build_dir = output_dir or fmt.report_root / "build"
    build_dir.mkdir(parents=True, exist_ok=True)
    markdown_path = build_dir / f"{OUTPUT_STEM}.md"
    csv_path = build_dir / f"{OUTPUT_STEM}.csv"
    output_docx = build_dir / f"{OUTPUT_STEM}.docx"
    stats = build_results_markdown(markdown_path, csv_path, OUTPUT_STEM)
    reference_docx = ensure_reference_docx(fmt)
    _run_pandoc(markdown_path, output_docx, reference_docx)
    postprocess_docx(output_docx, fmt)
    _set_landscape_a3(output_docx, fmt)
    map_stats = _postprocess_map_pages(output_docx, fmt)
    stats.update(map_stats)
    stats.update({
        "markdown": str(markdown_path),
        "csv": str(csv_path),
        "docx": str(output_docx),
    })
    return stats


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--format", type=Path, default=DEFAULT_FORMAT_PATH)
    parser.add_argument("--output-dir", type=Path, default=None)
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    stats = build_results_table_deliverable(
        format_path=args.format,
        output_dir=args.output_dir,
    )
    print(
        "Built results-table deliverable: "
        f"{stats['docx']} ({stats['selected_site_rows']} site rows, "
        f"{stats['copied_maps']} maps)"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
