# man_hours: 5.0
"""Apply report_format.json styling to a generated .docx."""

from __future__ import annotations

import re

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from report_format_config import ReportFormatConfig


_NUMERIC_PATTERN = re.compile(r"^\s*[+-]?\$?[0-9][0-9,]*(?:\.[0-9]+)?\s*%?\s*$")

_HEADING_STYLE_LEVEL = {
    "Heading 1": ("headings", "h1"),
    "Heading 2": ("headings", "h2"),
    "Heading 3": ("headings", "h3"),
    "Heading 4": ("headings", "h4"),
    "Title": ("title",),
    "Subtitle": ("subtitle",),
}


def _set_run_font(run, *, name: str, size_pt: float, bold: bool = False,
                  italic: bool = False) -> None:
    run.font.name = name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    rpr = run._element.get_or_add_rPr()
    for tag in ("rFonts",):
        existing = rpr.find(qn(f"w:{tag}"))
        if existing is not None:
            rpr.remove(existing)
    rfonts = OxmlElement("w:rFonts")
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:cs"), name)
    rfonts.set(qn("w:eastAsia"), name)
    rpr.append(rfonts)


def _style_spec(config: ReportFormatConfig, style_name: str) -> tuple[str, float, bool, bool, int]:
    if style_name in _HEADING_STYLE_LEVEL:
        keys = _HEADING_STYLE_LEVEL[style_name]
        node = config.data["typography"]
        for key in keys:
            node = node[key]
        weight = str(node.get("weight", "regular"))
        style = str(node.get("style", "regular"))
        return (
            str(node["font_family"]),
            float(node["size_pt"]),
            weight in ("bold", "semibold"),
            style == "italic",
            config.alignment(str(node["alignment"])),
        )
    body = config.data["typography"]["body"]
    return (
        str(body["font_family"]),
        float(body["size_pt"]),
        False,
        False,
        config.alignment(str(body["alignment"])),
    )


def _apply_paragraph_style(paragraph, config: ReportFormatConfig) -> None:
    style_name = paragraph.style.name if paragraph.style else ""
    if style_name.startswith("TOC"):
        return

    font_name, size_pt, bold, italic, alignment = _style_spec(config, style_name)
    pf = paragraph.paragraph_format
    pf.alignment = alignment
    if style_name in ("Normal",) or not style_name.startswith("Heading"):
        if style_name not in ("Title", "Subtitle") and not style_name.startswith("Heading"):
            pf.line_spacing = config.line_spacing("body")
            pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    for run in paragraph.runs:
        _set_run_font(run, name=font_name, size_pt=size_pt, bold=bold, italic=italic)


def _set_cell_vertical_center(cell) -> None:
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def _set_cell_borders(cell, *, vertical: bool) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    existing = tcPr.find(qn("w:tcBorders"))
    if existing is not None:
        tcPr.remove(existing)
    borders = OxmlElement("w:tcBorders")
    edges = ("top", "left", "bottom", "right") if vertical else ("top", "bottom")
    for edge in edges:
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), "000000")
        borders.append(border)
    tcPr.append(borders)


def _set_table_borders(table, *, vertical: bool) -> None:
    tblPr = table._tbl.tblPr
    existing = tblPr.find(qn("w:tblBorders"))
    if existing is not None:
        tblPr.remove(existing)
    borders = OxmlElement("w:tblBorders")
    edges = ["top", "bottom", "insideH"]
    if vertical:
        edges.extend(["left", "right", "insideV"])
    for edge in edges:
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), "000000")
        borders.append(border)
    tblPr.append(borders)


def _set_table_autofit(table) -> None:
    tblPr = table._tbl.tblPr
    for tag in ("w:tblW", "w:tblLayout"):
        existing = tblPr.find(qn(tag))
        if existing is not None:
            tblPr.remove(existing)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "autofit")
    tblPr.append(layout)


def _shade_cell(cell, hex_color: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    existing = tcPr.find(qn("w:shd"))
    if existing is not None:
        tcPr.remove(existing)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def _set_row_as_header(row) -> None:
    trPr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    trPr.append(header)


def _cell_is_numeric(text: str) -> bool:
    return bool(text) and bool(_NUMERIC_PATTERN.match(text.strip()))


def _column_alignments(table, config: ReportFormatConfig) -> list[WD_ALIGN_PARAGRAPH]:
    if len(table.rows) <= 1:
        return [config.alignment("left") for _ in range(len(table.columns))]

    alignments = [config.alignment("left")] * len(table.columns)
    for ci in range(len(table.columns)):
        numeric_cells = 0
        non_empty = 0
        for ri, row in enumerate(table.rows):
            if ri == 0:
                continue
            text = row.cells[ci].text.strip()
            if not text:
                continue
            non_empty += 1
            if _cell_is_numeric(text):
                numeric_cells += 1
        if non_empty and numeric_cells / non_empty >= 0.6:
            alignments[ci] = config.alignment(
                str(config.data["tables"]["cell_alignment"]["numeric"])
            )
    return alignments


def postprocess_docx(docx_path, config: ReportFormatConfig) -> dict[str, int]:
    doc = Document(docx_path)
    table_body = config.data["tables"]["body_row"]
    vertical_rules = config.table_vertical_rules()
    header_hex = config.header_shade_hex()
    table_spacing = config.table_line_spacing()

    body_paragraph_count = 0
    for paragraph in doc.paragraphs:
        style_name = paragraph.style.name if paragraph.style else ""
        if style_name.startswith("Heading") or style_name in ("Title", "TOC Heading"):
            _apply_paragraph_style(paragraph, config)
            continue
        if style_name.startswith("TOC"):
            continue
        _apply_paragraph_style(paragraph, config)
        body_paragraph_count += 1

    table_count = 0
    for table in doc.tables:
        _set_table_borders(table, vertical=vertical_rules)
        if config.table_layout_autofit():
            _set_table_autofit(table)
        alignments = _column_alignments(table, config)

        for ri, row in enumerate(table.rows):
            for ci, cell in enumerate(row.cells):
                _set_cell_vertical_center(cell)
                _set_cell_borders(cell, vertical=vertical_rules)
                alignment = alignments[ci] if ci < len(alignments) else config.alignment("left")
                for paragraph in cell.paragraphs:
                    pf = paragraph.paragraph_format
                    pf.alignment = alignment
                    pf.line_spacing = table_spacing
                    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
                    pf.space_before = Pt(0)
                    pf.space_after = Pt(0)
                    for run in paragraph.runs:
                        _set_run_font(
                            run,
                            name=str(table_body["font_family"]),
                            size_pt=float(table_body["size_pt"]),
                            bold=ri == 0,
                        )

                if ri == 0:
                    _shade_cell(cell, header_hex)

        if table.rows:
            _set_row_as_header(table.rows[0])
        table_count += 1

    doc.save(docx_path)
    return {"body_paragraphs": body_paragraph_count, "tables": table_count}
