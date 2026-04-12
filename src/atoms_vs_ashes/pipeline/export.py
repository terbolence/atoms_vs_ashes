# man_hours: 4.0
"""Export the entire database into a structured 4-sheet XLSX file.

Sheets:
1. **Sites** — One row per site with *all* data flattened: 46 GEM columns,
   Natural Hazards, Human Hazards, Radiological, Emergency Planning,
   Infrastructure, and aggregated Observations.
2. **Ownership** — Full ownership table joined with site name.
3. **Screening Verdicts** — Per-site per-SMR screening results.
4. **SMR Designs** — Reference table of reactor designs.
"""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import select
from sqlalchemy.orm import Session

from atoms_vs_ashes.db.models import (
    ScreeningVerdict,
    Site,
    SiteEmergencyPlanning,
    SiteHumanHazards,
    SiteInfrastructureV2,
    SiteNaturalHazards,
    SiteObservation,
    SiteOwnership,
    SiteRadiological,
    SiteUnit,
    SmrDesign,
)
from atoms_vs_ashes.logging import get_logger

log = get_logger(__name__)

_EXTENDED_KEYS = [
    "Unit name",
    "Conversion to (fuel)",
    "Conversion to (GEM unit ID)",
    "Alternate Fuel",
    "Major area (prefecture, district)",
    "Subregion",
    "Permit Parsed",
    "Captive",
    "Captive industry use",
    "Captive residential use",
    "CHP",
    "Capacity factor",
    "Plant age (years)",
    "Heat rate (Btu per kWh)",
    "Emission factor (kg of CO2 per TJ)",
    "Annual CO2 (million tonnes / annum)",
    "Remaining plant lifetime (years)",
    "Lifetime CO2 (million tonnes)",
    "China capacity payment recipient",
]

_SITE_COLUMNS = [
    "site_id",
    "gem_unit_phase_id",
    "gem_location_id",
    "country_code",
    "country_name",
    "wiki_url",
    "name",
    "alternative_names",
    "owner_operator",
    "owner_gem_id",
    "parent_company",
    "parent_gem_id",
    "installed_capacity_mw",
    "status",
    "start_year",
    "retired_year",
    "planned_retirement",
    "coal_phaseout_year",
    "net_zero_year",
    "combustion_technology",
    "coal_type",
    "coal_source",
    "location",
    "local_area",
    "subnational_unit",
    "region",
    "latitude",
    "longitude",
    "location_accuracy",
    "permits",
    "permit_date",
    "plant_type",
    "grid_voltage_kv",
    "grid_capacity_mw",
    "cooling_water_source",
    "site_area_ha",
    "elevation_m",
    "unit_count",
    "operating_capacity_mw",
]

_DOMAIN_MODELS: list[tuple[str, type]] = [
    ("nh", SiteNaturalHazards),
    ("hh", SiteHumanHazards),
    ("ri", SiteRadiological),
    ("ep", SiteEmergencyPlanning),
    ("ns", SiteInfrastructureV2),
]


def _autofit_columns(ws: Worksheet) -> None:
    """Set column widths to fit the header text (capped at 50)."""
    for col_idx, col_cells in enumerate(ws.iter_cols(min_row=1, max_row=1), start=1):
        header_value = col_cells[0].value
        width = min(max(len(str(header_value or "")) + 2, 10), 50)
        ws.column_dimensions[get_column_letter(col_idx)].width = width


def _freeze_header(ws: Worksheet) -> None:
    ws.freeze_panes = "A2"


def _domain_columns(model: type) -> list[str]:
    """Return column names from a domain model, excluding the join key."""
    return [c.key for c in model.__table__.columns if c.key != "site_id"]


def _domain_row(obj: object, model: type, prefix: str) -> dict:
    """Extract domain columns from an ORM object, prefixed to avoid clashes."""
    result: dict = {}
    for col_name in _domain_columns(model):
        val = getattr(obj, col_name, None)
        result[f"{prefix}_{col_name}"] = val
    return result


def _build_unit_detail_summary(units: list) -> str | None:
    """Build a human-readable summary of a plant's generating units."""
    if not units:
        return None
    parts = []
    for u in units:
        cap = f"{float(u.capacity_mw):.0f}MW" if u.capacity_mw else "?MW"
        name = u.unit_name or u.gem_unit_phase_id or "unit"
        parts.append(f"{name}: {cap} ({u.status or '?'})")
    return " | ".join(parts)


def _build_sites_df(session: Session) -> pd.DataFrame:
    """Build the flattened Sites sheet.

    Merges: 46 GEM columns + unit details + all domain-table columns
    (Natural Hazards, Human Hazards, Radiological, Emergency Planning,
    Infrastructure) + aggregated Observations.
    """
    sites = session.execute(select(Site)).scalars().all()

    all_units = session.execute(
        select(SiteUnit).order_by(SiteUnit.site_id, SiteUnit.unit_name)
    ).scalars().all()
    units_by_site: dict[str, list] = {}
    for u in all_units:
        units_by_site.setdefault(str(u.site_id), []).append(u)

    domain_lookup: dict[str, dict[str, object]] = {}
    for prefix, model in _DOMAIN_MODELS:
        objs = session.execute(select(model)).scalars().all()
        for obj in objs:
            domain_lookup.setdefault(str(obj.site_id), {})[prefix] = obj

    obs_stmt = select(SiteObservation).order_by(
        SiteObservation.site_id, SiteObservation.criterion_id
    )
    all_obs = session.execute(obs_stmt).scalars().all()
    obs_by_site: dict[str, list[str]] = {}
    for obs in all_obs:
        key = str(obs.site_id)
        entry = f"[{obs.criterion_id}] ({obs.impact or '?'}/{obs.confidence or '?'}) {obs.observation}"
        obs_by_site.setdefault(key, []).append(entry)

    rows: list[dict] = []
    for s in sites:
        row: dict = {}

        for col in _SITE_COLUMNS:
            val = getattr(s, col, None)
            if col == "alternative_names" and val:
                row[col] = "; ".join(val)
            else:
                row[col] = val

        ext = s.extended_data or {}
        for key in _EXTENDED_KEYS:
            row[key] = ext.get(key)

        sid = str(s.site_id)
        row["unit_details"] = _build_unit_detail_summary(units_by_site.get(sid, []))

        site_domains = domain_lookup.get(sid, {})
        for prefix, model in _DOMAIN_MODELS:
            obj = site_domains.get(prefix)
            if obj:
                row.update(_domain_row(obj, model, prefix))
            else:
                for col_name in _domain_columns(model):
                    row[f"{prefix}_{col_name}"] = None

        site_obs = obs_by_site.get(sid, [])
        row["observations"] = " | ".join(site_obs) if site_obs else None

        rows.append(row)

    return pd.DataFrame(rows)


def _build_ownership_df(session: Session) -> pd.DataFrame:
    stmt = (
        select(SiteOwnership, Site.name.label("site_name"), Site.country_name)
        .join(Site, SiteOwnership.site_id == Site.site_id)
    )
    results = session.execute(stmt).all()
    rows: list[dict] = []
    for own, site_name, country_name in results:
        row = {c.key: getattr(own, c.key) for c in SiteOwnership.__table__.columns}
        row["site_name"] = site_name
        row["country_name"] = country_name
        rows.append(row)
    return pd.DataFrame(rows)


def _build_verdicts_df(session: Session) -> pd.DataFrame:
    stmt = (
        select(
            ScreeningVerdict,
            Site.name.label("site_name"),
            SmrDesign.name.label("smr_name"),
        )
        .join(Site, ScreeningVerdict.site_id == Site.site_id)
        .join(SmrDesign, ScreeningVerdict.smr_key == SmrDesign.smr_key)
    )
    results = session.execute(stmt).all()
    rows: list[dict] = []
    for v, site_name, smr_name in results:
        row = {c.key: getattr(v, c.key) for c in ScreeningVerdict.__table__.columns}
        row["site_name"] = site_name
        row["smr_name"] = smr_name
        rows.append(row)
    return pd.DataFrame(rows)


def _build_smr_df(session: Session) -> pd.DataFrame:
    objs = session.execute(select(SmrDesign)).scalars().all()
    return pd.DataFrame(
        [{c.key: getattr(obj, c.key) for c in SmrDesign.__table__.columns} for obj in objs]
    )


def export_to_xlsx(session: Session, output_path: str | Path) -> Path:
    """Query all tables and write a 4-sheet XLSX workbook.

    Returns the resolved output path.
    """
    out = Path(output_path)
    out.parent.mkdir(parents=True, exist_ok=True)

    sheets: list[tuple[str, pd.DataFrame]] = [
        ("Sites", _build_sites_df(session)),
        ("Ownership", _build_ownership_df(session)),
        ("Screening Verdicts", _build_verdicts_df(session)),
        ("SMR Designs", _build_smr_df(session)),
    ]

    log.info("Exporting %d sheets to %s", len(sheets), out)

    with pd.ExcelWriter(out, engine="openpyxl") as writer:
        for sheet_name, df in sheets:
            if df.empty:
                df = pd.DataFrame({"(no data)": []})
            for col in df.columns:
                if pd.api.types.is_datetime64_any_dtype(df[col]):
                    df[col] = df[col].dt.tz_localize(None)
            df.to_excel(writer, sheet_name=sheet_name, index=False)
            ws: Worksheet = writer.sheets[sheet_name]
            _autofit_columns(ws)
            _freeze_header(ws)
            log.info("  Sheet '%s': %d rows, %d cols", sheet_name, len(df), len(df.columns))

    log.info("Export complete: %s", out)
    return out
