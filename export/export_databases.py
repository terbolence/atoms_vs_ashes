#!/usr/bin/env python3
"""Comprehensive database export orchestrator for atoms_vs_ashes.

Historically this file carried the full column catalog, every sheet
builder, and the CLI in one place (1 100+ lines). It is now the slim
orchestrator: builders live in :mod:`export._sheet_builders_core` and
:mod:`export._sheet_builders_merged`; column-clear-naming lives in
:mod:`export._column_mappings` (backed by ``export/column_mappings.csv``).

Features (unchanged from the previous monolith):

- Automatic versioning with timestamp.
- Graceful fallback for missing database columns.
- Optional clear column naming for external stakeholders.
- Excel formatting with auto-fit columns and frozen headers.
- Comprehensive summary reporting + column mapping reference sheet.
- Merged-DB-specific sheets: LLM verdict consensus, promoted
  observations, cell-level merge audit, source_db provenance summary,
  and the enrichment_runs ledger.
- Scoring artefact sheets: Ranking Scores + Composite Rankings (new
  in Phase 1.5 of the siting-report plan).

Usage:

    python export/export_databases.py               # api + llm (legacy default)
    python export/export_databases.py --database all   # api + llm + merged
    python export/export_databases.py --database merged
    python export/export_databases.py --clear-names
    python export/export_databases.py --version stakeholder_v1 -o /tmp/ex
"""

from __future__ import annotations

import os
import sys
from pathlib import Path

import click
import pandas as pd
from dotenv import load_dotenv
from sqlalchemy import text


def _repo_root() -> Path:
    """Find project root by looking for pyproject.toml."""
    here = Path(__file__).resolve()
    for d in (here.parent, *here.parents):
        if (d / "pyproject.toml").exists():
            return d
    return Path.cwd()


def _setup_path_and_env() -> None:
    """Add src to Python path and load environment."""
    root = _repo_root()
    src_path = root / "src"
    if str(src_path) not in sys.path:
        sys.path.insert(0, str(src_path))
    if str(root) not in sys.path:
        sys.path.insert(0, str(root))
    load_dotenv(root / ".env")


# We must set up sys.path BEFORE importing sibling helpers because
# ``export`` is a script directory, not an installed package.
_setup_path_and_env()

from export._column_mappings import COLUMN_MAPPINGS  # noqa: E402
from export._excel_utils import apply_clear_naming, sanitize_for_excel  # noqa: E402
from export._sheet_builders_core import (  # noqa: E402
    build_composite_rankings_df,
    build_ownership_df,
    build_ranking_scores_df,
    build_sites_df_safe,
    build_smr_df,
    build_standard_sites_df,
    build_verdicts_df,
)
from export._sheet_builders_merged import (  # noqa: E402
    build_enrichment_runs_df,
    build_llm_observations_df,
    build_llm_verdicts_df,
    build_merge_audit_df,
    build_provenance_summary_df,
)


_DB_NAME_BY_PROFILE: dict[str, str] = {
    "api": "atoms_vs_ashes",
    "llm": "atoms_vs_ashes_llm",
    "merged": "atoms_vs_ashes_merged",
}

_CORE_SHEETS = (
    ("Ownership", build_ownership_df),
    ("Screening Verdicts", build_verdicts_df),
    ("SMR Designs", build_smr_df),
    ("Ranking Scores", build_ranking_scores_df),
    ("Composite Rankings", build_composite_rankings_df),
)

_MERGED_SHEETS = (
    ("Provenance Summary", build_provenance_summary_df),
    ("LLM Verdicts", build_llm_verdicts_df),
    ("LLM Observations", build_llm_observations_df),
    ("Merge Audit", build_merge_audit_df),
    ("Enrichment Runs", build_enrichment_runs_df),
)


def _table_exists(session, table_name: str) -> bool:
    """Return True if ``table_name`` exists in the current DB."""
    try:
        nested = session.begin_nested()
        session.execute(text(f"SELECT 1 FROM {table_name} LIMIT 1"))
        nested.commit()
        return True
    except Exception:
        nested.rollback()
        return False


def _build_sites_sheet(session, use_clear_names: bool) -> pd.DataFrame:
    """Build the Sites sheet with a safe-mode fallback."""
    try:
        nested = session.begin_nested()
        df = build_standard_sites_df(session)
        nested.commit()
    except Exception as e:
        nested.rollback()
        click.echo(f"  ! Sites standard export failed, using safe method: {e}")
        nested2 = session.begin_nested()
        df = build_sites_df_safe(session)
        nested2.commit()
    return apply_clear_naming(df, "Sites", use_clear_names)


def _append_sheet(sheets, session, name, fn):
    """Run ``fn(session)`` inside a savepoint and append to ``sheets``."""
    try:
        nested = session.begin_nested()
        df = fn(session)
        nested.commit()
        sheets.append((name, df))
        click.echo(f"  ✓ {name}: {len(df)} rows, {len(df.columns)} columns")
    except Exception as e:
        nested.rollback()
        click.echo(f"  ✗ {name} failed: {e}")
        sheets.append((name, pd.DataFrame({"error": [f"{name} export failed: {e}"]})))


def _write_excel(output_path: Path, sheets: list[tuple[str, pd.DataFrame]]) -> None:
    """Write ``sheets`` to ``output_path`` with formatting niceties."""
    click.echo(f"Writing Excel file to {output_path}")
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        for sheet_name, df in sheets:
            if df.empty:
                df = pd.DataFrame({"(no data)": []})
            df = sanitize_for_excel(df)
            df.to_excel(writer, sheet_name=sheet_name, index=False)
            try:
                from openpyxl.utils import get_column_letter

                ws = writer.sheets[sheet_name]
                for col_idx, col_cells in enumerate(
                    ws.iter_cols(min_row=1, max_row=1), start=1
                ):
                    header_value = col_cells[0].value
                    width = min(max(len(str(header_value or "")) + 2, 10), 50)
                    ws.column_dimensions[get_column_letter(col_idx)].width = width
                ws.freeze_panes = "A2"
            except Exception:
                pass


def export_database(
    db_profile: str, output_path: Path, use_clear_names: bool = False
) -> Path | None:
    """Export a single database profile to Excel."""
    from sqlalchemy import create_engine
    from sqlalchemy.orm import sessionmaker as _sessionmaker

    from atoms_vs_ashes.config import Settings

    db_name = _DB_NAME_BY_PROFILE.get(db_profile)
    if db_name is None:
        click.echo(f"✗ Unknown database profile: {db_profile!r}", err=True)
        return None
    os.environ["POSTGRES_DB"] = db_name
    settings = Settings()

    engine = create_engine(settings.database.url, echo=False, pool_pre_ping=True)
    SessionLocal = _sessionmaker(bind=engine)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    try:
        session = SessionLocal()
        try:
            click.echo(f"Exporting {db_profile.upper()} database...")
            sheets: list[tuple[str, pd.DataFrame]] = []

            if _table_exists(session, "sites"):
                try:
                    sites_df = _build_sites_sheet(session, use_clear_names)
                    sheets.append(("Sites", sites_df))
                    click.echo(
                        f"  ✓ Sites: {len(sites_df)} rows, "
                        f"{len(sites_df.columns)} columns"
                    )
                except Exception as e:
                    click.echo(f"  ✗ Sites failed completely: {e}")
                    sheets.append(
                        ("Sites", pd.DataFrame({"error": [f"Sites export failed: {e}"]}))
                    )

            for name, fn in _CORE_SHEETS:
                _append_sheet(sheets, session, name, fn)

            if db_profile == "merged":
                for name, fn in _MERGED_SHEETS:
                    _append_sheet(sheets, session, name, fn)

            if use_clear_names:
                mapping_rows = [
                    {
                        "Original_Column": old,
                        "Clear_Column_Name": new,
                        "Description": desc,
                    }
                    for old, (new, desc) in COLUMN_MAPPINGS.items()
                ]
                if mapping_rows:
                    sheets.append(("Column Mappings", pd.DataFrame(mapping_rows)))
                    click.echo(f"  ✓ Column Mappings: {len(mapping_rows)} mappings")

            _write_excel(output_path, sheets)
            click.echo(f"✓ {db_profile.upper()} export complete: {output_path}")
            return output_path
        finally:
            session.close()
            engine.dispose()
    except Exception as e:
        click.echo(f"✗ {db_profile.upper()} export failed: {e}", err=True)
        return None


if __name__ == "__main__":
    from export._cli import main

    main()
