# man_hours: 2.5
"""Tests for report_format.json loading and reference template generation."""

from __future__ import annotations

from pathlib import Path

import pytest

docx = pytest.importorskip("docx")
Document = docx.Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

REPO_ROOT = Path(__file__).resolve().parent.parent
SCRIPTS = REPO_ROOT / "scripts"
FORMAT_PATH = (
    REPO_ROOT
    / "report/version 1.02/output/report/writing plan/report_format.json"
)


@pytest.fixture
def scripts_on_path(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.syspath_prepend(str(SCRIPTS))


def test_format_config_resolves_report_root(scripts_on_path: None) -> None:
    from report_format_config import ReportFormatConfig

    fmt = ReportFormatConfig.load(FORMAT_PATH)
    assert fmt.report_root.name == "report"
    assert fmt.chapters_dir.exists()
    assert fmt.build_path("reference_docx").parent.name == "build"
    assert fmt.data["tables"]["borders"]["vertical_rules"] is True
    assert fmt.data["tables"]["borders"]["cell_all_sides"] is True


def test_reference_docx_uses_times_new_roman_and_justified_body(
    scripts_on_path: None,
    tmp_path: Path,
) -> None:
    from make_reference_docx import build_reference_docx
    from report_format_config import ReportFormatConfig

    fmt = ReportFormatConfig.load(FORMAT_PATH)
    out = tmp_path / "reference.docx"
    build_reference_docx(out, fmt)

    doc = Document(out)
    normal = doc.styles["Normal"]
    assert normal.font.name == "Times New Roman"
    assert normal.font.size.pt == 11
    assert normal.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY

    heading1 = doc.styles["Heading 1"]
    assert heading1.font.name == "Times New Roman"
    assert heading1.font.size.pt == 16

    footer = doc.sections[0].footer.paragraphs[0]
    assert footer.alignment == WD_ALIGN_PARAGRAPH.CENTER


def test_ensure_reference_docx_writes_hash_and_rebuilds_on_change(
    scripts_on_path: None,
    tmp_path: Path,
) -> None:
    from report_format_config import ReportFormatConfig, ensure_reference_docx

    fmt_json = tmp_path / "report_format.json"
    fmt_json.write_text(
        ReportFormatConfig.load(FORMAT_PATH).path.read_text(encoding="utf-8"),
        encoding="utf-8",
    )
    fmt = ReportFormatConfig.load(fmt_json)
    fmt.data["build"] = {
        "reference_docx": "build/reference.docx",
        "output_docx": "build/out.docx",
        "merged_markdown": "build/merged.md",
    }
    fmt_json.write_text(
        __import__("json").dumps(fmt.data, indent=2) + "\n",
        encoding="utf-8",
    )
    fmt = ReportFormatConfig.load(fmt_json)

    first = ensure_reference_docx(fmt)
    first_hash = fmt.reference_docx_hash_path().read_text(encoding="utf-8").strip()
    second = ensure_reference_docx(fmt)
    assert second == first
    assert fmt.reference_docx_hash_path().read_text(encoding="utf-8").strip() == first_hash

    fmt.data["typography"]["body"]["size_pt"] = 10.5
    fmt_json.write_text(
        __import__("json").dumps(fmt.data, indent=2) + "\n",
        encoding="utf-8",
    )
    fmt = ReportFormatConfig.load(fmt_json)
    rebuilt = ensure_reference_docx(fmt)
    assert rebuilt == first
    assert fmt.reference_docx_hash_path().read_text(encoding="utf-8").strip() != first_hash

    doc = Document(rebuilt)
    assert doc.styles["Normal"].font.size.pt == 10.5


def test_writing_controls_do_not_contain_em_dash() -> None:
    controls = [
        FORMAT_PATH,
        FORMAT_PATH.parent / "writingStyle.md",
        FORMAT_PATH.parent / "writingDecisions.md",
        FORMAT_PATH.parent / "v1_2_iteration_controls.md",
        FORMAT_PATH.parent / "prompts/specialists/writing_quality_auditor.md",
    ]
    for path in controls:
        assert "\u2014" not in path.read_text(encoding="utf-8"), path
